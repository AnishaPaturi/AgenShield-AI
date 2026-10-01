"""
build_college_document.py
Generates the complete Project Stage-I documentation in docs/college/PS - Document Format.docx
with strict, 100% uniform alignment and indentation across all sections, front matter, tables,
figures, and body paragraphs.
"""

import shutil
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Add backend directory to sys.path
sys.path.insert(0, str(Path(__file__).parent))
from write_chapters import append_all_chapters

DOC_PATH = Path("docs/college/PS - Document Format.docx")
BACKUP_PATH = Path("docs/college/PS - Document Format_backup.docx")
IMG_DIR = Path("docs/college/extracted_images")

# Always reload from pristine backup if it exists
if BACKUP_PATH.exists():
    shutil.copyfile(BACKUP_PATH, DOC_PATH)
    print(f"Restored pristine template from {BACKUP_PATH}")
else:
    shutil.copyfile(DOC_PATH, BACKUP_PATH)
    print(f"Backed up original document to {BACKUP_PATH}")

doc = docx.Document(DOC_PATH)

def format_run(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)

def set_para_bounds(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_in=0.0, right_in=0.0, first_in=0.0, before_pt=0, after_pt=6, line_sp=1.15):
    p.alignment = align
    pf = p.paragraph_format
    pf.left_indent = Inches(left_in)
    pf.right_indent = Inches(right_in)
    pf.first_line_indent = Inches(first_in)
    pf.space_before = Pt(before_pt)
    pf.space_after = Pt(after_pt)
    pf.line_spacing = line_sp

print("Step 1: Fixing Front Matter with Uniform Alignments...")

# 1. Cover page title
p_title = doc.paragraphs[3]
set_para_bounds(p_title, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=12, after_pt=12)
p_title.text = "AGENTSHIELD AI: AN AUTONOMOUS MULTI-AGENT FRAMEWORK FOR MULTI-CLOUD INFRASTRUCTURE-AS-CODE SECURITY"
for r in p_title.runs:
    format_run(r, font_name="Times New Roman", size_pt=16, bold=True)

# 2. Cover page guide
p_guide = doc.paragraphs[25]
set_para_bounds(p_guide, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=6, after_pt=6)
p_guide.text = "Ms. K. Vishal Reddy, Assistant Professor, Department of CSE"
for r in p_guide.runs:
    format_run(r, font_name="Times New Roman", size_pt=12, bold=True)

# 3. Certificate supervisor name
for p in doc.paragraphs[68:75]:
    if "(Ms/Mr XXXX)" in p.text:
        p.text = p.text.replace("(Ms/Mr XXXX)", "(Ms. K. Vishal Reddy)")
    set_para_bounds(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_in=0, right_in=0, first_in=0, after_pt=6)

# 4. Declaration
p_dec_heading = doc.paragraphs[155]
set_para_bounds(p_dec_heading, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=18, after_pt=12)

p_dec_body = doc.paragraphs[158]
set_para_bounds(p_dec_body, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_in=0, right_in=0, first_in=0, after_pt=12)

# Clear old placeholders
for i in [168, 169, 170]:
    doc.paragraphs[i].text = ""

def format_student_table_block(p_hdr, p_names):
    # Set clean tab stop at 4.2 inches
    for p in [p_hdr, p_names]:
        set_para_bounds(p, align=WD_ALIGN_PARAGRAPH.LEFT, left_in=1.0, right_in=0, first_in=0, after_pt=4)
        pPr = p._p.get_or_add_pPr()
        old_tabs = pPr.find(qn('w:tabs'))
        if old_tabs is not None:
            pPr.remove(old_tabs)
        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'left')
        tab.set(qn('w:pos'), '6048')  # 4.2 inches in dxa
        tabs.append(tab)
        pPr.append(tabs)

    p_hdr.text = "Student Name\tRoll no."
    for r in p_hdr.runs:
        format_run(r, font_name="Times New Roman", size_pt=12, bold=True)

    p_names.text = (
        "Anisha Paturi\t23BD1A050E\n"
        "Ch Parinamika Bhanu\t23BD1A0518\n"
        "Ch Venkata Vahini\t23BD1A051D\n"
        "Sravani Janak\t23BD1A051Y"
    )
    for r in p_names.runs:
        format_run(r, font_name="Times New Roman", size_pt=12, bold=True)

format_student_table_block(doc.paragraphs[165], doc.paragraphs[167])

# 5. Acknowledgement
p_ack_heading = doc.paragraphs[171]
set_para_bounds(p_ack_heading, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=18, after_pt=12)

for p in doc.paragraphs[173:188]:
    if p.text.strip():
        if "Ms. XXXX" in p.text:
            p.text = p.text.replace("Ms. XXXX", "Ms. K. Vishal Reddy")
        set_para_bounds(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_in=0, right_in=0, first_in=0, after_pt=6)

# Clear 191-194 and put clean block at 189
for i in [191, 192, 193, 194]:
    doc.paragraphs[i].text = ""
format_student_table_block(doc.paragraphs[189], doc.paragraphs[190])

# 6. Abstract
p_abs_heading = doc.paragraphs[198]
set_para_bounds(p_abs_heading, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=18, after_pt=12)

abstract_text = (
    "Cloud-native enterprises increasingly rely on Infrastructure-as-Code (IaC) paradigms—such as "
    "Terraform, AWS CloudFormation, Kubernetes manifests, and Helm charts—to automate the provisioning "
    "and orchestration of heterogeneous multi-cloud topologies. However, architectural security misconfigurations "
    "codified at the template layer (such as unrestricted security group ingress rules, unencrypted object "
    "stores, wildcard IAM privilege grants, and hardcoded credentials) propagate instantly into production environments, "
    "accounting for over 73% of cloud enterprise security breaches. Traditional static analysis tools (e.g., Checkov, tfsec, KICS) "
    "suffer from severe syntactic myopia and generate substantial false-positive cascades (32%–48%) due to their inability "
    "to resolve cross-module variable flows and complex resource dependency graphs. Furthermore, existing scanners operate strictly "
    "as diagnostic instruments without automated remediation, resulting in an industry-average Mean Time to Remediation (MTTR) "
    "exceeding 24 days. To address these critical challenges, this project presents AgentShield AI, an autonomous, stateful "
    "multi-agent DevSecOps framework designed for multi-cloud syntactic verification, calibrated secret interception, "
    "consensus-driven dual-LLM remediation, and sandbox-verified patch synthesis. AgentShield AI employs eight specialized "
    "autonomous agents orchestrated over a typed LangGraph execution context graph (Γ). The framework incorporates a polyglot "
    "AST parser dispatcher, a dual-engine secret interceptor combining pattern signatures with Shannon entropy analysis, a hybrid "
    "Retrieval-Augmented Generation (RAG) engine fusing dense Qdrant vector embeddings with sparse BM25 lexical scores via "
    "Reciprocal Rank Fusion (RRF) over CIS and NIST compliance standards, and an LLM consensus engine evaluating cross-model agreement "
    "between frontier models. To guarantee that generated remediations are safe and syntactically sound, candidate patches undergo a "
    "two-tier validation harness combining static linting with simulated LocalStack sandbox runtime dry-runs, equipped with automated "
    "self-healing retry loops. When evaluated on benchmark templates, AgentShield AI delivers high-precision vulnerability detection, "
    "verified Unified Git Diff code patches, attack-path blast radius prioritization, and automated compliance reporting (JSON, Markdown, "
    "HTML, SARIF, and PDF), drastically reducing developer remediation fatigue and elevating cloud security posture."
)
p_abs_body = doc.paragraphs[201]
set_para_bounds(p_abs_body, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_in=0, right_in=0, first_in=0, after_pt=6, line_sp=1.15)
p_abs_body.text = abstract_text
for r in p_abs_body.runs:
    format_run(r, font_name="Times New Roman", size_pt=12, bold=False)

print("Step 2: Formatting List of Figures Table with Uniform Columns...")

p_lof_heading = doc.paragraphs[202]
set_para_bounds(p_lof_heading, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=18, after_pt=12)

figures_list = [
    ("Figure 1.1", "End-to-End Multi-Agent Architecture Flow of AgentShield AI", "4"),
    ("Figure 1.2", "Six-Stage Subsystem Architecture & Execution Trace", "5"),
    ("Figure 2.1", "Base Research Article in IEEE Access (Toprani & Madisetti, 2025)", "7"),
    ("Figure 2.2", "Base Paper Objective, Architecture, and Performance Benchmark", "8"),
    ("Figure 2.3", "Comparative Analysis of State-of-the-Art Approaches in Literature", "9"),
    ("Figure 2.4", "Existing System Architecture Workflow and Limitations", "10"),
    ("Figure 2.5", "Advantages of the Existing System", "11"),
    ("Figure 2.6", "Architectural Disadvantages and Bottlenecks of Existing Systems", "12"),
    ("Figure 2.7", "Proposed System Objective and Key Pipeline Architecture", "13"),
    ("Figure 2.8", "Architectural Advantages and Dimensional Comparison of Proposed Framework", "14"),
    ("Figure 2.9", "Trade-off and Disadvantage Analysis of Proposed Multi-Agent Approach", "15"),
    ("Figure 3.1", "Software Requirements Specification Functional & Agent Architecture", "18"),
    ("Figure 3.2", "Comprehensive Functional, Non-Functional, and Hardware/Software Requirements", "20"),
    ("Figure 3.3", "High-Level End-to-End System Workflow and Agent State Lifecycle", "23"),
    ("Figure 4.1", "UML Use Case Diagram for AgentShield AI Platform", "27"),
    ("Figure 4.2", "UML Class Diagram & Pydantic Data Contracts Architecture", "29"),
    ("Figure 4.3", "UML Sequence Diagram - Execution Trace from Ingestion to Sandbox Remediation", "32"),
    ("Figure 4.4", "UML Statechart Diagram - Workspace & Remediation Patch Lifecycle", "35"),
    ("Figure 4.5", "UML Deployment Diagram - Containerized Microservices and Cloud Infrastructure", "38")
]

table1 = doc.tables[1]
table1.alignment = WD_TABLE_ALIGNMENT.CENTER
while len(table1.rows) < len(figures_list) + 1:
    table1.add_row()

for idx, (fig_no, fig_title, page_no) in enumerate(figures_list):
    row_cells = table1.rows[idx + 1].cells
    row_cells[0].text = f"{idx + 1}."
    row_cells[1].text = f"{fig_no}: {fig_title}"
    row_cells[2].text = page_no
    for c_idx, cell in enumerate(row_cells):
        for p in cell.paragraphs:
            p.paragraph_format.left_indent = Inches(0)
            p.paragraph_format.right_indent = Inches(0)
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx == 0 or c_idx == 2) else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                format_run(r, font_name="Times New Roman", size_pt=11, bold=False)

print("Step 3: Formatting Table of Contents with Precise Dot Leaders...")

p_contents = doc.paragraphs[206]
set_para_bounds(p_contents, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=18, after_pt=12)

def set_toc_paragraph(p, text, is_chapter=False, tab_pos_dxa="9600", has_leader=True):
    pPr = p._p.get_or_add_pPr()
    old_tabs = pPr.find(qn('w:tabs'))
    if old_tabs is not None:
        pPr.remove(old_tabs)
    
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    if is_chapter:
        pf.left_indent = Inches(0)
        pf.first_line_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.space_before = Pt(4)
        pf.space_after = Pt(2)
    else:
        pf.left_indent = Inches(0.25)
        pf.first_line_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
    
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot' if has_leader else 'none')
    tab.set(qn('w:pos'), str(tab_pos_dxa))
    tabs.append(tab)
    pPr.append(tabs)
    
    p.text = text
    for r in p.runs:
        format_run(r, font_name="Times New Roman", size_pt=11, bold=is_chapter)

# Format header DESCRIPTION PAGE
p_desc = doc.paragraphs[208]
set_toc_paragraph(p_desc, "DESCRIPTION\tPAGE", is_chapter=True, tab_pos_dxa="9600", has_leader=False)

toc_entries = [
    ("CHAPTER - 1", "1", True),
    ("INTRODUCTION", "2-5", True),
    ("1.1 Purpose of the Project", "2", False),
    ("1.2 Problem with Existing Systems", "2", False),
    ("1.3 Proposed System", "3", False),
    ("1.4 Scope of the Project", "4", False),
    ("1.5 Architecture Diagram", "4", False),
    ("CHAPTER - 2", "6", True),
    ("LITERATURE SURVEY", "7-16", True),
    ("2.1 Overview & Evolution of IaC Security", "7", False),
    ("2.2 Review of Base Paper (Toprani & Madisetti)", "7", False),
    ("2.3 Traditional Rule-Based Static Analyzers", "9", False),
    ("2.4 CSPM and Runtime Scanners", "10", False),
    ("2.5 Machine Learning & Neural Approaches", "11", False),
    ("2.6 Secret Detection Mechanisms", "12", False),
    ("2.7 Existing System Analysis", "13", False),
    ("2.8 Proposed System Innovations", "14", False),
    ("2.9 Trade-off and Disadvantage Analysis", "15", False),
    ("CHAPTER - 3", "17", True),
    ("SOFTWARE REQUIREMENT SPECIFICATION", "18-24", True),
    ("3.1 Introduction to SRS", "18", False),
    ("3.2 Role and Purpose of SRS", "18", False),
    ("3.3 Requirements Specification Document", "19", False),
    ("3.4 Functional Requirements", "20", False),
    ("3.5 Non-Functional Requirements", "21", False),
    ("3.6 Performance Requirements", "22", False),
    ("3.7 Software Requirements", "23", False),
    ("3.8 Hardware Requirements", "23", False),
    ("CHAPTER - 4", "25", True),
    ("SYSTEM DESIGN", "26-42", True),
    ("4.1 Introduction to UML", "26", False),
    ("4.2 UML Diagrams in AgentShield AI", "26", False),
    ("4.3 Use Case Diagram", "27", False),
    ("4.4 Class Diagram", "29", False),
    ("4.5 Sequence Diagram", "32", False),
    ("4.6 State Chart Diagram", "35", False),
    ("4.7 Deployment Diagram", "38", False),
    ("4.8 TECHNOLOGIES USED", "40", True),
]

# Ensure we have enough paragraphs between 209 and the section break
curr_p_idx = 209
for title, page, is_chap in toc_entries:
    p = doc.paragraphs[curr_p_idx]
    set_toc_paragraph(p, f"{title}\t{page}", is_chapter=is_chap, tab_pos_dxa="9600", has_leader=True)
    curr_p_idx += 1

print(f"Step 4: Cleaning up empty paragraphs between TOC and Chapter 1 (from index {curr_p_idx} to 289)...")
# Delete all extra paragraphs between curr_p_idx and paragraph 289 (which contains sectPr)
paragraphs_to_remove = 289 - curr_p_idx
for _ in range(paragraphs_to_remove):
    p = doc.paragraphs[curr_p_idx]
    p._element.getparent().remove(p._element)

# Now doc.paragraphs[curr_p_idx] is CHAPTER-1 with sectPr, and curr_p_idx+1 is INTRODUCTION
p_ch1 = doc.paragraphs[curr_p_idx]
set_para_bounds(p_ch1, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=20, after_pt=6)
p_ch1.text = "CHAPTER - 1"
p_ch1.style = "Heading 1"
for r in p_ch1.runs:
    format_run(r, font_name="Times New Roman", size_pt=16, bold=True)

p_intro = doc.paragraphs[curr_p_idx + 1]
set_para_bounds(p_intro, align=WD_ALIGN_PARAGRAPH.CENTER, left_in=0, right_in=0, first_in=0, before_pt=6, after_pt=12)
p_intro.text = "INTRODUCTION"
p_intro.style = "Heading 1"
for r in p_intro.runs:
    format_run(r, font_name="Times New Roman", size_pt=16, bold=True)

# Truncate any stray template paragraphs that existed after curr_p_idx + 1
while len(doc.paragraphs) > curr_p_idx + 2:
    p = doc.paragraphs[-1]
    p._element.getparent().remove(p._element)

print("Step 5: Updating College Running Footers...")
for s_idx in range(len(doc.sections)):
    s = doc.sections[s_idx]
    for p in s.footer.paragraphs:
        for node in p._p.xpath('.//w:t'):
            if '|  Project' in node.text:
                node.text = node.text.replace('|  Project', '|  AgentShield AI')
            elif node.text.strip() == 'Name':
                node.text = ''

print("Step 6: Writing Chapters 1, 2, 3, and 4 with Unified Styles...")
append_all_chapters(doc)

print("Step 7: Universal Document Alignment Normalization Pass...")
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    pf = p.paragraph_format
    
    # Check if image paragraph
    is_img = any('graphic' in r._r.xml for r in p.runs)
    if is_img:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.left_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.first_line_indent = Inches(0)
        continue
    
    # Check if caption
    if txt.startswith("Figure ") or txt.startswith("Table "):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.left_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.first_line_indent = Inches(0)
        continue

    # Skip front matter and Table of Contents paragraphs (handled by specific formatting)
    if idx < curr_p_idx:
        continue

    # Chapter headings: CHAPTER - 1, INTRODUCTION, etc.
    if txt.startswith("CHAPTER -") or txt in [
        "INTRODUCTION", "LITERATURE SURVEY", 
        "SOFTWARE REQUIREMENT SPECIFICATION", "SYSTEM DESIGN"
    ]:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.left_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.first_line_indent = Inches(0)
        continue

    # Section Headings (e.g., 1.1, 1.2, 2.1, 3.1, 4.1, 4.8)
    if any(txt.startswith(f"{ch}.") for ch in ["1", "2", "3", "4"]) and not txt.startswith("Figure"):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.first_line_indent = Inches(0)
        continue

    # If list item (starts with bullet or number like "1. ", "• ")
    if txt.startswith("•") or (len(txt) > 2 and txt[0].isdigit() and txt[1] in [".", ")"]):
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.left_indent = Inches(0.25)
        pf.right_indent = Inches(0)
        pf.first_line_indent = Inches(-0.25)
        pf.line_spacing = 1.15
        continue

    # Regular body paragraphs
    if txt:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.left_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.first_line_indent = Inches(0)
        pf.line_spacing = 1.15

print("Step 8: Saving completed document...")
doc.save(DOC_PATH)
print(f"Successfully saved cleanly formatted document to {DOC_PATH}")
