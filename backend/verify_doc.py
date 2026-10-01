import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document('docs/college/PS - Document Format.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}, Total sections: {len(doc.sections)}")

# Check key front matter paragraphs
print("\n--- FRONT MATTER ALIGNMENT CHECK ---")
for idx in [3, 25, 155, 158, 165, 167, 171, 180, 189, 198, 201, 202, 206, 208, 209, 210, 211]:
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        pf = p.paragraph_format
        print(f"[{idx:3d}] align={str(p.alignment):18s} left={str(pf.left_indent):10s} first={str(pf.first_line_indent):10s} | {p.text[:45]!r}")

# Check Chapter 1 transition
print("\n--- CHAPTER 1 & 2 TRANSITION CHECK ---")
for i, p in enumerate(doc.paragraphs):
    if p.text.startswith("CHAPTER -"):
        pf = p.paragraph_format
        has_sect = 'sectPr' in p._p.xml
        print(f"[{i:3d}] align={str(p.alignment):18s} left={str(pf.left_indent):10s} sectPr={has_sect} | {p.text!r}")
        next_p = doc.paragraphs[i+1]
        print(f"      next title: align={str(next_p.alignment):18s} left={str(next_p.paragraph_format.left_indent):10s} | {next_p.text!r}")

# Sample body paragraphs, lists, headings, and figures
print("\n--- SAMPLE CONTENT PARAGRAPHS CHECK ---")
counts = {"JUSTIFY": 0, "CENTER": 0, "LEFT": 0, "OTHER": 0}
for p in doc.paragraphs[247:]:
    t = p.text.strip()
    if not t:
        continue
    if p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
        counts["JUSTIFY"] += 1
    elif p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
        counts["CENTER"] += 1
    elif p.alignment == WD_ALIGN_PARAGRAPH.LEFT:
        counts["LEFT"] += 1
    else:
        counts["OTHER"] += 1

print(f"Paragraph alignment distribution in Chapters 1-4: {counts}")

# Check TOC tab stops
print("\n--- TOC TAB STOP CHECK ---")
for i in [208, 209, 210, 211, 212]:
    p = doc.paragraphs[i]
    tabs_xml = p._p.xpath('.//w:tabs/w:tab')
    tab_info = [(t.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'), 
                 t.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}leader'), 
                 t.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pos')) for t in tabs_xml]
    print(f"[{i}] text={p.text[:30]!r} | tabs={tab_info}")
