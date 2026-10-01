import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document('docs/college/PS - Document Format.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")

print("\n--- FRONT MATTER (0 to 238) NON-EMPTY PARAGRAPHS ---")
for i in range(0, 239):
    p = doc.paragraphs[i]
    t = p.text.strip()
    if t:
        pf = p.paragraph_format
        print(f"[{i:3d}] style={p.style.name:15s} align={str(p.alignment):18s} left={str(pf.left_indent):10s} first={str(pf.first_line_indent):10s} | {t[:60]}")
