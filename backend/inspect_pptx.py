import pptx
from pathlib import Path

prs = pptx.Presentation("docs/college/Review-AgentSheild-AI.pptx")
print("Total Slides:", len(prs.slides))

out_dir = Path("docs/college/extracted_images")
out_dir.mkdir(exist_ok=True)

img_index = 0
for i, slide in enumerate(prs.slides):
    slide_title = ""
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                texts.append(txt)
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            img = shape.image
            img_index += 1
            ext = img.ext
            filename = f"slide_{i+1}_img_{img_index}.{ext}"
            file_path = out_dir / filename
            with open(file_path, "wb") as f:
                f.write(img.blob)
            print(f"Slide {i+1} saved image {filename} ({len(img.blob)} bytes)")

    title = texts[0] if texts else "No title"
    print(f"--- Slide {i+1}: {title[:60]} ---")
    for t in texts[:3]:
        print("   *", t[:80].replace("\n", " "))
