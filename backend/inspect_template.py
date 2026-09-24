import docx
import os

template_path = r"C:\Users\anish\OneDrive\College\project-clg\AgenShield-AI\docs\paper\Paper-Template-IMPACT-2027.docx"
doc = docx.Document(template_path)

print(f"Loaded template successfully. Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")

for i in range(12):
    p = doc.paragraphs[i]
    print(f"P{i} [{p.style.name}]: text='{p.text[:60]}'")
    for j, r in enumerate(p.runs):
        print(f"   r{j}: text='{r.text}' font={r.font.name} size={r.font.size.pt if r.font.size else None} bold={r.bold} italic={r.italic}")
