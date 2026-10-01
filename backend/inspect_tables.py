import docx

doc = docx.Document('docs/college/PS - Document Format_backup.docx')
print("Total tables in backup:", len(doc.tables))
for i, t in enumerate(doc.tables):
    parent = t._tbl.getparent()
    print(f"Table {i}: rows={len(t.rows)}, cols={len(t.columns)}")
    # Find preceding paragraph
    prev = t._tbl.getprevious()
    if prev is not None:
        print(f"  Preceding element tag: {prev.tag}")
        # print text if paragraph
        if prev.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(prev, doc)
            print(f"  Preceding paragraph text: {p.text!r}")
