import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document('docs/college/PS - Document Format.docx')
print("Inspecting paragraph alignments and indents across the document:")
for i in [3, 25, 68, 167, 180, 201, 206, 209, 210, 211, 235, 236, 237, 238, 239, 240, 241, 242, 243, 244, 245]:
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        pf = p.paragraph_format
        print(f"[{i:3d}] style={p.style.name:15s} align={str(p.alignment):25s} left={str(pf.left_indent):10s} first={str(pf.first_line_indent):10s} text={p.text[:40]!r}")
