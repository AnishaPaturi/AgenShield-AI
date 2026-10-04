"""
apply_footer.py
Applies the required institutional running footer to docs/college/PS_-_Document_Format_Formatted.docx:
Left: Keshav Memorial Institute of Technology (KMIT)
Center: Dynamic Page Number (Page No)
Right: AGENTSHIELD AI: AN AUTONOMOUS MULTI-AGENT FRAMEWORK FOR MULTI-CLOUD INFRASTRUCTURE-AS-CODE SECURITY
"""

import os
import shutil
from pathlib import Path
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

DOC_PATH = Path(r"docs/college/PS_-_Document_Format_Formatted.docx")
BACKUP_PATH = Path(r"docs/college/PS_-_Document_Format_Formatted_pre_footer.docx")

# Ensure backup exists
if not BACKUP_PATH.exists():
    shutil.copyfile(DOC_PATH, BACKUP_PATH)
    print(f"Created safety backup at {BACKUP_PATH}")

doc = docx.Document(DOC_PATH)

def build_footer_table(footer, width_in=6.52):
    """
    Constructs a 1-row, 3-column borderless table with a top divider line
    and perfectly formatted runs for Left, Center (Page No), and Right.
    """
    # Clear any existing paragraphs/tables in this footer
    for p in list(footer.paragraphs):
        p._p.getparent().remove(p._p)
    for t in list(footer.tables):
        t._tbl.getparent().remove(t._tbl)
        
    table = footer.add_table(rows=1, cols=3, width=Inches(width_in))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Table border XML: crisp top border line, no other borders
    tblBorders = parse_xml(
        '<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '</w:tblBorders>'
    )
    table._tbl.tblPr.append(tblBorders)
    
    # Set cell widths
    row = table.rows[0]
    row.cells[0].width = Inches(2.30)
    row.cells[1].width = Inches(0.60)
    row.cells[2].width = Inches(3.62)
    
    # Set cell top/bottom margins
    tcPr_list = [row.cells[0]._tc.get_or_add_tcPr(),
                 row.cells[1]._tc.get_or_add_tcPr(),
                 row.cells[2]._tc.get_or_add_tcPr()]
    for tcPr in tcPr_list:
        tcMar = parse_xml(
            '<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:top w:w="80" w:type="dxa"/>'
            '<w:bottom w:w="40" w:type="dxa"/>'
            '<w:left w:w="0" w:type="dxa"/>'
            '<w:right w:w="0" w:type="dxa"/>'
            '</w:tcMar>'
        )
        tcPr.append(tcMar)
        
    # --- Cell 0: Left ---
    p0 = row.cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p0.paragraph_format.space_before = Pt(3)
    p0.paragraph_format.space_after = Pt(0)
    p0.paragraph_format.line_spacing = 1.0
    r0 = p0.add_run("Keshav Memorial Institute of Technology (KMIT)")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(8.5)
    r0._r.get_or_add_rPr().append(parse_xml(
        '<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
    ))
    
    # --- Cell 1: Center (Dynamic Page Number) ---
    p1 = row.cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(3)
    p1.paragraph_format.space_after = Pt(0)
    p1.paragraph_format.line_spacing = 1.0
    
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    rf = OxmlElement('w:r')
    rfPr = OxmlElement('w:rPr')
    rfPr.append(parse_xml(
        '<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
    ))
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '18') # 9pt
    rfPr.append(sz)
    rf.append(rfPr)
    tf = OxmlElement('w:t')
    tf.text = '1'
    rf.append(tf)
    fld.append(rf)
    p1._p.append(fld)
    
    # --- Cell 2: Right ---
    p2 = row.cells[2].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_before = Pt(3)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.0
    r2 = p2.add_run("AGENTSHIELD AI: AN AUTONOMOUS MULTI-AGENT FRAMEWORK FOR MULTI-CLOUD INFRASTRUCTURE-AS-CODE SECURITY")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(7.5)
    r2._r.get_or_add_rPr().append(parse_xml(
        '<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
    ))

print("Applying footer to all document sections...")

# Section 0 controls footer1.xml (Sections 0 to 6 inherit from it)
build_footer_table(doc.sections[0].footer)

# Section 7 controls footer2.xml (Sections 7 and 8 inherit from it)
build_footer_table(doc.sections[7].footer)

# Also ensure any other unlinked section footers receive the exact same table
for idx, sec in enumerate(doc.sections):
    if not sec.footer.is_linked_to_previous and idx not in [0, 7]:
        build_footer_table(sec.footer)
        print(f"Applied to unlinked section {idx}")

# Also clean up the redundant blank paragraphs on cover page (Section 0)
# so 'DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING' sits neatly on page 1
# rather than spilling over onto page 2
print("Optimizing cover page paragraph spacing...")
blank_indices = []
for i in range(27, 41):
    if doc.paragraphs[i].text.strip() == '':
        blank_indices.append(i)

# Remove 5 redundant blank paragraphs to allow clean 1-page cover
if len(blank_indices) >= 8:
    for idx in reversed(blank_indices[:5]):
        p = doc.paragraphs[idx]
        p._p.getparent().remove(p._p)
    print(f"Removed 5 redundant blank lines on cover page. Remaining blank lines: {len(blank_indices) - 5}")

doc.save(DOC_PATH)
print(f"Successfully saved updated document to {DOC_PATH}")
